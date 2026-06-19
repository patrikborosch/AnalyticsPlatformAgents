#!/usr/bin/env python3
"""
Weekly Status PPTX template -- Fabric Skills
=================================================================
Copy this file to a working directory of your choice (e.g. a
gitignored .local/weekly/ folder), rename it to
generate-weekly-{YYYY-MM-DD}.py, then fill in the blocks marked
"EDIT ME" below.

Run (PowerShell):
    pip install python-pptx
    python generate-weekly-{YYYY-MM-DD}.py

IMPORTANT: close PowerPoint before regenerating (file lock):
    Get-Process POWERPNT -ErrorAction SilentlyContinue |
        ForEach-Object { $_.CloseMainWindow() | Out-Null }

Slide order:
    1. Title
    2. Milestones & KPIs (8 cards + 5 KPI stack)
    3. Test Infrastructure (8 cards + 5 KPI stack) -- feature-focused
    4. Skill PRs Merged This Week (table)
    5. Open Skill PRs -- Review Status (table)
    6. ADO Status (table with clickable links)
    7. Next Week
=================================================================
"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ADO = "https://powerbi.visualstudio.com/Trident/_workitems/edit"

# -------- Design tokens (vivid text + soft backgrounds) ----------
C = {
    # Vivid text + accents
    "navy":      "1E2761",
    "darkText":  "1E293B",
    "muted":     "64748B",
    "accent":    "3B82F6",
    "green":     "16A34A",
    "amber":     "D97706",
    "red":       "DC2626",
    "teal":      "0891B2",
    # Muted fills + chrome
    "navyFill":  "3F5485",
    "offWhite":  "F7F9FC",
    "white":     "FFFFFF",
    "ice":       "E3EAF5",
    "tableHead": "5E6F94",
    "tableRow1": "FAFBFD",
    "tableRow2": "F1F5FB",
    # Status / delta pill fills
    "greenFill":  "7FB39A",
    "amberFill":  "D9B36A",
    "redFill":    "C58585",
    "tealFill":   "8FB8C4",
    "accentFill": "7DA0D6",
}
FONT_HEADER = "Georgia"
FONT_BODY = "Calibri"

STATUS = {
    "notStarted": ("Not Started", C["amberFill"]),
    "inProgress": ("In Progress", C["greenFill"]),
    "completed":  ("Completed",   C["greenFill"]),
    "blocked":    ("Blocked",     C["redFill"]),
}

KPI_Y = [0.95, 1.85, 2.75, 3.65, 4.55]

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


# -------- Low-level helpers --------------------------------------

def _set_slide_bg(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_color)


def _zero_text_margins(tf):
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True


def _add_shadow(shape):
    """Inject a soft outer shadow into the shape's spPr."""
    sp = shape._element
    spPr = sp.find(qn("p:spPr"))
    if spPr is None:
        return
    existing = spPr.find(qn("a:effectLst"))
    if existing is not None:
        spPr.remove(existing)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:outerShdw blurRad="50800" dist="25400" dir="8100000" rotWithShape="0">'
        '<a:srgbClr val="000000"><a:alpha val="8000"/></a:srgbClr>'
        '</a:outerShdw>'
        '</a:effectLst>'
    )
    elem = etree.fromstring(xml)
    # effectLst goes after fill/ln but before scene3d/sp3d/extLst (per CT_ShapeProperties).
    anchor = None
    for tag in ("a:scene3d", "a:sp3d", "a:extLst"):
        found = spPr.find(qn(tag))
        if found is not None:
            anchor = found
            break
    if anchor is not None:
        anchor.addprevious(elem)
    else:
        spPr.append(elem)


def add_rect(slide, x, y, w, h, *, fill=None, line=None, line_width=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line is not None:
        shp.line.color.rgb = RGBColor.from_string(line)
        if line_width is not None:
            shp.line.width = Pt(line_width)
    else:
        shp.line.fill.background()
    if shadow:
        _add_shadow(shp)
    _zero_text_margins(shp.text_frame)
    return shp


def add_oval(slide, x, y, w, h, *, fill=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
    shp.line.fill.background()
    _zero_text_margins(shp.text_frame)
    return shp


def add_text(slide, text, x, y, w, h, *,
             font_size=11, font_face=FONT_BODY, color=None,
             bold=False, italic=False, align="left", valign="top"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    _zero_text_margins(tf)
    tf.vertical_anchor = _VALIGN.get(valign, MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = _ALIGN.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.name = font_face
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    return tb


def _fill_cell(cell, text, opts):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(opts.get("fill", C["white"]))
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = _VALIGN.get(opts.get("valign", "middle"), MSO_ANCHOR.MIDDLE)
    tf = cell.text_frame
    tf.word_wrap = True
    # text_frame has one default empty paragraph; reset its content.
    p = tf.paragraphs[0]
    p.alignment = _ALIGN.get(opts.get("align", "left"), PP_ALIGN.LEFT)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run()
    run.text = text
    run.font.size = Pt(opts.get("font_size", 10))
    run.font.name = opts.get("font_face", FONT_BODY)
    if opts.get("color") is not None:
        run.font.color.rgb = RGBColor.from_string(opts["color"])
    run.font.bold = opts.get("bold", False)
    run.font.italic = opts.get("italic", False)
    if opts.get("hyperlink"):
        run.hyperlink.address = opts["hyperlink"]


def _set_cell_borders(cell, color="D1D5DB", width_pt=0.5):
    """Add solid borders on all 4 sides; insert before fill per CT_TableCellProperties order."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    w_emu = int(width_pt * 12700)
    line_tags = {qn(f"a:{e}") for e in ("lnL", "lnR", "lnT", "lnB", "lnTlBr", "lnBlTr")}
    # Anchor: first non-line child (e.g. solidFill) so lines stay schema-ordered.
    anchor = None
    for child in tcPr:
        if child.tag not in line_tags:
            anchor = child
            break
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        existing = tcPr.find(qn(f"a:{edge}"))
        if existing is not None:
            tcPr.remove(existing)
        xml = (
            f'<a:{edge} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'w="{w_emu}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:prstDash val="solid"/><a:round/>'
            f'</a:{edge}>'
        )
        elem = etree.fromstring(xml)
        if anchor is not None:
            anchor.addprevious(elem)
        else:
            tcPr.append(elem)


def add_table(slide, x, y, *, col_widths, row_h, header, body, border_color="D1D5DB"):
    rows_n = 1 + len(body)
    cols_n = len(col_widths)
    w_total = sum(col_widths)
    h_total = row_h * rows_n
    graphic = slide.shapes.add_table(
        rows_n, cols_n,
        Inches(x), Inches(y), Inches(w_total), Inches(h_total),
    )
    table = graphic.table

    # Strip the default table style (built-in firstRow / banding) so cell fills win cleanly.
    tbl = table._tbl
    tblPr = tbl.tblPr
    for tag in ("a:tableStyleId",):
        node = tblPr.find(qn(tag))
        if node is not None:
            tblPr.remove(node)
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    tblPr.set("firstCol", "0")
    tblPr.set("bandCol", "0")
    tblPr.set("lastRow", "0")
    tblPr.set("lastCol", "0")

    for i, cw in enumerate(col_widths):
        table.columns[i].width = Inches(cw)
    for i in range(rows_n):
        table.rows[i].height = Inches(row_h)

    rows_all = [header] + list(body)
    for r, row in enumerate(rows_all):
        for c, (text, opts) in enumerate(row):
            cell = table.cell(r, c)
            _fill_cell(cell, text, opts)
            if border_color:
                _set_cell_borders(cell, color=border_color, width_pt=0.5)
    return table


# -------- Card + KPI block (shared by SLIDE 2 and SLIDE 3) -------

def render_cards_and_kpis(slide, cards, kpis):
    """8 milestone-style cards on the left, 5 KPI tiles on the right."""
    for i, m in enumerate(cards):
        y_top = 0.95 + i * 0.55
        card_h = 0.48

        # Card body + left status-color accent bar
        add_rect(slide, 0.6, y_top, 6.30, card_h, fill=C["white"], shadow=True)
        add_rect(slide, 0.6, y_top, 0.07, card_h, fill=m["color"])

        # Number circle
        add_oval(slide, 0.86, y_top + 0.09, 0.30, 0.30, fill=C["navyFill"])
        add_text(slide, m["num"], 0.86, y_top + 0.09, 0.30, 0.30,
                 font_size=11, color=C["white"], bold=True,
                 align="center", valign="middle")

        # Delta badge (DONE / NEW) -- small, theme-tinted, secondary to status pill
        if m.get("delta"):
            delta_color = C["greenFill"] if m["delta"] == "DONE" else C["accentFill"]
            label = "DONE" if m["delta"] == "DONE" else "NEW"
            add_rect(slide, 4.20, y_top + 0.13, 0.55, 0.22, fill=delta_color)
            add_text(slide, label, 4.20, y_top + 0.13, 0.55, 0.22,
                     font_size=7, color=C["white"], bold=True,
                     align="center", valign="middle")

        # Title + one-line desc
        add_text(slide, m["title"], 1.28, y_top + 0.00, 2.62, 0.24,
                 font_size=11, color=C["darkText"], bold=True)
        add_text(slide, m["desc"], 1.28, y_top + 0.23, 2.62, 0.24,
                 font_size=8, color=C["muted"])

        # Status pill (filled)
        add_rect(slide, 4.80, y_top + 0.08, 0.95, 0.32, fill=m["color"])
        add_text(slide, m["label"], 4.80, y_top + 0.08, 0.95, 0.32,
                 font_size=9, color=C["white"], bold=True,
                 align="center", valign="middle")

        # ETA pill (outlined, navy text)
        add_rect(slide, 5.82, y_top + 0.08, 1.03, 0.32,
                 fill=C["ice"], line=C["tableHead"], line_width=0.5)
        add_text(slide, m["eta"], 5.82, y_top + 0.08, 1.03, 0.32,
                 font_size=9, color=C["navy"], bold=True,
                 align="center", valign="middle")

    for k in kpis:
        add_rect(slide, 7.0, k["y"], 2.5, 0.85, fill=C["white"], shadow=True)
        add_text(slide, k["big"], 7.0, k["y"] + 0.02, 2.5, 0.42,
                 font_size=k["big_fs"], color=k["big_color"], bold=True,
                 align="center", valign="middle")
        add_text(slide, k["label"], 7.0, k["y"] + 0.44, 2.5, 0.20,
                 font_size=9, color=C["darkText"], bold=True,
                 align="center", valign="middle")
        add_text(slide, k["sub"], 7.0, k["y"] + 0.63, 2.5, 0.20,
                 font_size=7, color=C["muted"], italic=True,
                 align="center", valign="middle")


# -------- Build deck ---------------------------------------------

def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]  # built-in blank layout
    # EDIT ME -- core properties
    prs.core_properties.author = "Fabric Skills Team"
    prs.core_properties.title = "Fabric Skills -- Weekly Summary ({MMM D - MMM D, YYYY})"

    # ============================================================
    # SLIDE 1 -- Title
    # ============================================================
    s1 = prs.slides.add_slide(blank)
    _set_slide_bg(s1, C["navyFill"])
    add_rect(s1, 0, 0, 10, 0.06, fill=C["accentFill"])

    add_text(s1, "Fabric Skills", 0.8, 1.4, 8.4, 1.0,
             font_size=48, font_face=FONT_HEADER, color=C["white"], bold=True)
    # EDIT ME -- date range
    add_text(s1, "Weekly Summary -- {Month D - Month D, YYYY}", 0.8, 2.5, 8.4, 0.6,
             font_size=22, color=C["ice"])
    # EDIT ME -- 1-line top-line highlights
    add_text(s1, "Highlights:  {release(s) shipped}  -  {N} PRs merged  -  {N} new skills landed",
             0.8, 3.4, 8.4, 0.45,
             font_size=14, color=C["muted"], italic=True)
    add_rect(s1, 0, 5.565, 10, 0.06, fill=C["accentFill"])

    # ============================================================
    # SLIDE 2 -- Milestones & KPIs
    # ============================================================
    s2 = prs.slides.add_slide(blank)
    _set_slide_bg(s2, C["offWhite"])
    add_text(s2, "Milestones & KPIs", 0.6, 0.3, 8.8, 0.6,
             font_size=28, font_face=FONT_HEADER, color=C["navy"], bold=True)

    # EDIT ME -- 5 KPIs (skill PRs reviewed/open, merged this week, total public skills,
    #                    skills awaiting public, latest public release)
    kpis = [
        {"y": KPI_Y[0], "big": "X / Y",  "big_fs": 28, "big_color": C["amber"],
         "label": "Skill PRs Reviewed",      "sub": "Z not ready (drafts + failing validation)"},
        {"y": KPI_Y[1], "big": "N",      "big_fs": 32, "big_color": C["green"],
         "label": "PRs Merged (this week)",  "sub": "{a} skill content -  {b} release / eval / CI"},
        {"y": KPI_Y[2], "big": "N",      "big_fs": 32, "big_color": C["teal"],
         "label": "Total Public Skills",     "sub": "+{n} new this week ({version})"},
        {"y": KPI_Y[3], "big": "N",      "big_fs": 32, "big_color": C["green"],
         "label": "Skills Awaiting Public",  "sub": "{caught-up note or backlog count}"},
        {"y": KPI_Y[4], "big": "v0.x.y", "big_fs": 24, "big_color": C["accent"],
         "label": "Latest Public Release",   "sub": "Shipped {date} -  prior {version} {date}"},
    ]

    # EDIT ME -- 8 milestone cards in board priority order.
    # delta: "DONE" -> completed this week  -  "NEW" -> new since last week  -  "" -> carryover
    def _ms(num, delta, title, desc, eta, status_key):
        label, color = STATUS[status_key]
        return {"num": num, "delta": delta, "title": title, "desc": desc,
                "eta": eta, "label": label, "color": color}

    milestones = [
        _ms("1", "DONE", "{Milestone 1 title}", "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("2", "DONE", "{Milestone 2 title}", "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("3", "",     "{Milestone 3 title}", "{1-line desc -  supporting PRs -  owner}", "ETA {date}", "inProgress"),
        _ms("4", "NEW",  "{Milestone 4 title}", "{1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("5", "",     "{Milestone 5 title}", "{1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("6", "",     "{Milestone 6 title}", "{1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("7", "",     "{Milestone 7 title}", "{1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("8", "",     "{Milestone 8 title}", "{1-line desc -  owner}", "ETA {date}", "inProgress"),
    ]
    render_cards_and_kpis(s2, milestones, kpis)

    # ============================================================
    # SLIDE 3 -- Test Infrastructure
    # ============================================================
    s3 = prs.slides.add_slide(blank)
    _set_slide_bg(s3, C["offWhite"])
    add_text(s3, "Test Infrastructure", 0.6, 0.3, 8.8, 0.6,
             font_size=28, font_face=FONT_HEADER, color=C["navy"], bold=True)
    # EDIT ME -- 1-line subtitle scoping the slide
    add_text(s3,
             "{eval framework -  smoke harness -  release automation -- what shipped this week and what's next}",
             0.6, 0.78, 8.8, 0.18,
             font_size=10, color=C["muted"], italic=True)

    # EDIT ME -- 8 infra cards (5 shipped this week + 3 in-flight, typical mix)
    infra_cards = [
        _ms("1", "",    "{Shipped feature 1}",   "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("2", "",    "{Shipped feature 2}",   "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("3", "",    "{Shipped feature 3}",   "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("4", "",    "{Shipped feature 4}",   "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("5", "",    "{Shipped feature 5}",   "{1-line desc -  owner}", "Shipped {date}", "completed"),
        _ms("6", "NEW", "{In-flight feature 1}", "DRAFT -  {1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("7", "NEW", "{In-flight feature 2}", "DRAFT -  {1-line desc -  owner}", "ETA {date}", "inProgress"),
        _ms("8", "NEW", "{In-flight feature 3}", "DRAFT -  {1-line desc -  owner}", "ETA {date}", "inProgress"),
    ]
    # EDIT ME -- 5 infra KPIs (features shipped / in flight / harness phase / smoke failures / next drop)
    infra_kpis = [
        {"y": KPI_Y[0], "big": "N",      "big_fs": 32, "big_color": C["green"],
         "label": "Features Shipped",   "sub": "{short list}"},
        {"y": KPI_Y[1], "big": "N",      "big_fs": 32, "big_color": C["accent"],
         "label": "Features In Flight", "sub": "DRAFTs opened this week"},
        {"y": KPI_Y[2], "big": "{phase}","big_fs": 22, "big_color": C["teal"],
         "label": "{Major eval surface}","sub": "{status note}"},
        {"y": KPI_Y[3], "big": "N",      "big_fs": 32, "big_color": C["green"],
         "label": "Smoke Run Failures", "sub": "{environment note}"},
        {"y": KPI_Y[4], "big": "{date}", "big_fs": 24, "big_color": C["accent"],
         "label": "Next Drop",          "sub": "{theme}"},
    ]
    render_cards_and_kpis(s3, infra_cards, infra_kpis)

    # ============================================================
    # SLIDE 4 -- PRs Merged This Week
    # ============================================================
    s4 = prs.slides.add_slide(blank)
    _set_slide_bg(s4, C["offWhite"])
    add_text(s4, "PRs Merged This Week", 0.6, 0.3, 8.8, 0.6,
             font_size=28, font_face=FONT_HEADER, color=C["navy"], bold=True)
    # EDIT ME -- subtitle counts
    add_text(s4,
             "{total} PRs merged  -  {skill_count} skill content ({new} new -  {upd} updated)  -  {infra_count} eval / CI / release infrastructure",
             0.6, 0.85, 8.8, 0.35,
             font_size=13, color=C["muted"])

    hdr = {"fill": C["tableHead"], "color": C["white"], "bold": True,
           "font_size": 11, "align": "left", "valign": "middle"}

    def cell(bg, color=C["darkText"], **kw):
        opts = {"fill": bg, "color": color, "font_size": 10,
                "align": "left", "valign": "middle"}
        opts.update(kw)
        return opts

    # EDIT ME -- merged PR rows. Category color hints:
    #   Skill=teal -  Release=accent -  Eval=accent -  CI=amber -  Infra=green
    merged_rows = [
        [
            ("#NNN", cell(C["tableRow1"])),
            ("{skill-name} -- NEW skill", cell(C["tableRow1"])),
            ("Skill", cell(C["tableRow1"], color=C["teal"], bold=True)),
            ("{Author}", cell(C["tableRow1"])),
        ],
        # Add more rows... alternate cell(C["tableRow1"]) / cell(C["tableRow2"]).
    ]
    add_table(
        s4, 0.2, 1.4,
        col_widths=[0.85, 5.45, 1.3, 2.0],
        row_h=0.27,
        header=[("PR", hdr), ("Title", hdr), ("Category", hdr), ("Author", hdr)],
        body=merged_rows,
    )

    # ============================================================
    # SLIDE 5 -- Open Skill PRs -- Review Status
    # ============================================================
    s5 = prs.slides.add_slide(blank)
    _set_slide_bg(s5, C["offWhite"])
    add_text(s5, "Open Skill PRs -- Review Status", 0.6, 0.3, 8.8, 0.6,
             font_size=28, font_face=FONT_HEADER, color=C["navy"], bold=True)
    # EDIT ME -- header summary
    add_text(s5,
             "Reviewed X of Y  -  Not ready: Z (drafts + failing validation)  -  Awaiting review: N  -  M new arrivals this week",
             0.6, 0.85, 8.8, 0.35,
             font_size=12, color=C["muted"])

    # EDIT ME -- open PR rows. Status color hints:
    #   Reviewed=green -  New This Week=teal -  Needs Review=amber
    open_rows = [
        [
            ("#NNN", cell(C["tableRow1"])),
            ("{Open PR title}", cell(C["tableRow1"])),
            ("{Author}", cell(C["tableRow1"])),
            ("Needs Review", cell(C["tableRow1"], color=C["amber"], bold=True)),
        ],
        # Add more rows... ~10-12 rows fit comfortably.
    ]
    add_table(
        s5, 0.2, 1.3,
        col_widths=[0.7, 5.0, 1.8, 1.3],
        row_h=0.30,
        header=[("PR#", hdr), ("Title", hdr), ("Author", hdr), ("Status", hdr)],
        body=open_rows,
    )
    add_text(s5, "+N more open PRs ({comma-separated short labels})",
             0.6, 5.0, 8.8, 0.4,
             font_size=9, color=C["muted"], italic=True)

    # ============================================================
    # SLIDE 6 -- ADO Status
    # ============================================================
    s6 = prs.slides.add_slide(blank)
    _set_slide_bg(s6, C["offWhite"])
    add_text(s6, "ADO Status", 0.6, 0.3, 8.8, 0.6,
             font_size=28, font_face=FONT_HEADER, color=C["navy"], bold=True)

    ado_hdr = {"fill": C["tableHead"], "color": C["white"], "bold": True,
               "font_size": 11, "align": "left", "valign": "middle"}

    def ado_cell(bg, color=C["darkText"], **kw):
        opts = {"fill": bg, "color": color, "font_size": 10,
                "align": "left", "valign": "middle"}
        opts.update(kw)
        return opts

    def ado_link(bg, url, **kw):
        return ado_cell(bg, color=C["accent"], hyperlink=url, **kw)

    def ado_status(bg, color):
        return {"fill": bg, "color": color, "bold": True,
                "font_size": 9, "align": "center", "valign": "middle"}

    # EDIT ME -- Active User Stories with tag `skills-for-fabric`. Skip Closed and New.
    #            ETA color: amber=this month -  teal=next month+.
    ado_rows = [
        [
            ("1", ado_cell(C["tableRow1"], bold=True)),
            ("{User Story title}", ado_link(C["tableRow1"], f"{ADO}/{{ID}}", bold=True)),
            ("In Progress", ado_status(C["tableRow1"], C["green"])),
            ("{Owner}", ado_cell(C["tableRow1"])),
            ("{M/D}", ado_cell(C["tableRow1"], color=C["amber"], bold=True)),
        ],
        # Optional sub-row for supporting child tasks / merged PRs
        [
            ("", ado_cell(C["tableRow2"])),
            ("  -> {supporting child task or merged PR list}", ado_cell(C["tableRow2"], color=C["green"])),
            ("", ado_cell(C["tableRow2"])),
            ("", ado_cell(C["tableRow2"])),
            ("", ado_cell(C["tableRow2"])),
        ],
        # Add more US rows...
    ]
    add_table(
        s6, 0.3, 1.0,
        col_widths=[0.35, 5.45, 1.1, 0.9, 1.05],
        row_h=0.32,
        header=[("#", ado_hdr), ("Item", ado_hdr), ("Status", ado_hdr),
                ("Owner", ado_hdr), ("ETA", ado_hdr)],
        body=ado_rows,
    )
    add_text(s6,
             "User Story titles are clickable ADO links  -  Tag filter: skills-for-fabric  -  Active only  -  Ordered by board priority",
             0.6, 4.3, 8.8, 0.3,
             font_size=9, color=C["muted"])
    add_text(s6,
             "Closed this week: {US title} ({ID})  -  {US title} ({ID})",
             0.6, 4.6, 8.8, 0.3,
             font_size=9, color=C["green"], italic=True)

    # ============================================================
    # SLIDE 7 -- Next Week
    # ============================================================
    s7 = prs.slides.add_slide(blank)
    _set_slide_bg(s7, C["navyFill"])
    add_rect(s7, 0, 0, 10, 0.06, fill=C["accentFill"])
    add_text(s7, "Next Week", 0.8, 0.5, 8.4, 0.7,
             font_size=30, font_face=FONT_HEADER, color=C["white"], bold=True)

    # EDIT ME -- 4-5 short focus items
    next_items = [
        ("1", "{focus item 1 -- concrete, short}"),
        ("2", "{focus item 2}"),
        ("3", "{focus item 3}"),
        ("4", "{focus item 4}"),
        ("5", "{focus item 5}"),
    ]
    for i, (icon, text) in enumerate(next_items):
        y_pos = 1.5 + i * 0.75
        add_oval(s7, 0.8, y_pos, 0.5, 0.5, fill=C["accentFill"])
        add_text(s7, icon, 0.8, y_pos, 0.5, 0.5,
                 font_size=18, color=C["white"], bold=True,
                 align="center", valign="middle")
        add_text(s7, text, 1.5, y_pos, 7.7, 0.5,
                 font_size=16, color=C["ice"], valign="middle")
    add_rect(s7, 0, 5.565, 10, 0.06, fill=C["accentFill"])

    # ---- Write file ------------------------------------------------
    # EDIT ME -- output file name
    out_path = Path(__file__).with_name("Fabric-Skills-Weekly-{YYYY-MM-DD}.pptx")
    prs.save(out_path)
    print(f"Created: {out_path}")


if __name__ == "__main__":
    build()
